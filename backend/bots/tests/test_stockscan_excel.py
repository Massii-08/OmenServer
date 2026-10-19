"""Scanner d'actions — le classeur Excel (écrit puis relu avec openpyxl)."""
from datetime import date

import pytest
from openpyxl import load_workbook

from backend.bots.stockscan import excel_out as xo


def _results(lang="fr"):
    return {
        "version": 1,
        "generated_at": "2026-10-17T07:45:12+00:00",
        "config": {"markets": ["italy", "switzerland"], "min_cap_meur": 1000, "min_value_meur": 5,
                   "window_h": 48, "top_n": 25, "shortlist": 40, "exclude_red": False, "lang": lang},
        "source": "screener",
        "markets": [{"id": "italy", "ok": True, "rows": 200, "error": None},
                    {"id": "switzerland", "ok": False, "rows": 0, "error": "HTTP 403"}],
        "counts": {"universe": 180, "shortlist": 40, "kept": 2, "excluded_no_catalyst": 30,
                   "excluded_red": 1},
        "warnings": [{"code": "market_failed", "market": "switzerland", "error": "HTTP 403"},
                     {"code": "degraded_columns", "columns": ["gap"]}],
        "rows": [
            {"rank": 1, "tv_symbol": "MIL:ENI", "symbol": "ENI.MI", "ticker": "ENI",
             "name": "Eni S.p.A.", "market": "italy", "currency": "EUR", "price": 15.48,
             "change_pct": 2.31, "rel_volume": 2.4, "score": 72, "market_cap_meur": 50200,
             "avg_value_meur": 201.24, "earnings_next": "2026-10-24",
             "flags": ["earnings_soon", "new"],
             "reasons": [{"code": "volume", "value": 2.4}, {"code": "news", "value": 3, "hours": 48},
                         {"code": "earnings", "value": "2026-10-24"}],
             "news": [{"title": "Eni beats estimates", "url": "https://www.tradingview.com/news/1",
                       "sentiment": "pos"}],
             "links": {"chart": "https://www.tradingview.com/chart/?symbol=MIL%3AENI"}},
            {"rank": 2, "tv_symbol": "MIL:STLAM", "symbol": "STLAM.MI", "ticker": "STLAM",
             "name": "=HYPERLINK(\"http://evil\")", "market": "italy", "currency": "EUR",
             "price": 9.12, "change_pct": -6.5, "rel_volume": None, "score": 15,
             "flags": ["red", "late"], "reasons": [],
             "news": [{"title": "+cmd|' /C calc'!A0", "url": "javascript:alert(1)", "sentiment": "neg"}],
             "links": {"chart": None}},
        ],
    }


def _open(tmp_path, lang="fr"):
    path = xo.write_workbook(_results(lang), tmp_path)
    return path, load_workbook(str(path))


def test_file_name_and_sheets(tmp_path):
    path, wb = _open(tmp_path)
    assert path.name == "scanner_actions_2026-10-17.xlsx"
    assert wb.sheetnames == ["Candidats", "Infos"]
    assert not list(tmp_path.glob(".*.tmp"))           # écriture atomique propre


@pytest.mark.parametrize("lang,sheet,first_header,file_prefix", [
    ("en", "Candidates", "Rank", "stock_scanner_"),
    ("it", "Candidati", "Pos.", "scanner_azioni_"),
])
def test_languages(tmp_path, lang, sheet, first_header, file_prefix):
    path, wb = _open(tmp_path, lang)
    assert path.name.startswith(file_prefix)
    assert wb.sheetnames[0] == sheet
    assert wb[sheet].cell(row=xo.DATA_START_ROW - 1, column=1).value == first_header


def test_candidate_rows(tmp_path):
    _, wb = _open(tmp_path)
    ws = wb["Candidats"]
    r = xo.DATA_START_ROW
    assert ws.cell(row=r, column=2).value == "ENI"
    assert ws.cell(row=r, column=4).value == "Italie (Borsa Italiana)"
    assert ws.cell(row=r, column=7).value == 2.31
    assert ws.cell(row=r, column=7).number_format == xo.FMT_PCT
    assert ws.cell(row=r, column=9).value == 72
    assert ws.cell(row=r, column=10).value == "résultats proches, nouveau"
    assert ws.cell(row=r, column=11).value == ("volume ×2,4 vs son marché · 3 news en 48 h"
                                               " · résultats le 24/10/2026")
    assert ws.cell(row=r, column=12).hyperlink.target == "https://www.tradingview.com/news/1"
    assert ws.cell(row=r, column=14).value.date() == date(2026, 10, 24)
    assert ws.cell(row=r, column=17).hyperlink.target.endswith("MIL%3AENI")
    assert ws.freeze_panes == "A%d" % xo.DATA_START_ROW
    assert ws.auto_filter.ref.startswith("A%d:" % (xo.DATA_START_ROW - 1))


def test_formula_injection_and_bad_links_are_neutralised(tmp_path):
    _, wb = _open(tmp_path)
    ws = wb["Candidats"]
    r = xo.DATA_START_ROW + 1
    assert ws.cell(row=r, column=3).value.startswith("'=")
    assert ws.cell(row=r, column=12).value.startswith("'+")
    assert ws.cell(row=r, column=12).hyperlink is None    # javascript: refusé
    assert ws.cell(row=r, column=17).value is None         # pas de lien, pas de texte
    assert ws.cell(row=r, column=8).value is None          # valeur absente = cellule vide


def test_info_sheet_lists_warnings_and_markets(tmp_path):
    _, wb = _open(tmp_path)
    texts = [c.value for row in wb["Infos"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "Marché en échec : switzerland (HTTP 403)" in texts
    assert "Colonnes du screener vides : gap" in texts
    assert any(t.startswith("échec") for t in texts)
    assert "Italie (Borsa Italiana), Suisse (SIX)" in texts


def test_reason_and_flag_texts():
    assert xo.reason_text({"code": "aligned", "value": 1.5}, "fr") == "le cours confirme (1,5 %)"
    assert xo.reason_text({"code": "aligned", "value": 1.5}, "en") == "price confirms (1.5 %)"
    assert xo.reason_text({"code": "earnings", "value": "2026-10-24"}, "en") == "earnings on 2026-10-24"
    assert xo.reason_text({"code": "mystery"}, "fr") == "mystery"
    assert xo.flag_text("late", "it") == "già partito"
    # chaque code de raison/drapeau/avertissement existe dans les trois langues
    for lang in ("fr", "en", "it"):
        assert set(xo.LABELS[lang]["reasons"]) == set(xo.LABELS["fr"]["reasons"])
        assert set(xo.LABELS[lang]["flags"]) == set(xo.LABELS["fr"]["flags"])
        assert set(xo.LABELS[lang]["warnings"]) == set(xo.LABELS["fr"]["warnings"])
        assert len(xo.LABELS[lang]["headers"]) == len(xo.WIDTHS)


def test_volume_reason_says_it_is_compared_to_its_market():
    """La valeur de la raison est le volume rapporté à la normale de SA place à
    cette heure (``rel_volume_adj``), pas le volume relatif brut."""
    reason = {"code": "volume", "value": 2.4}
    assert xo.reason_text(reason, "fr") == "volume ×2,4 vs son marché"
    assert xo.reason_text(reason, "en") == "volume ×2.4 vs its market"
    assert xo.reason_text(reason, "it") == "volume ×2,4 rispetto al suo mercato"


def test_rel_volume_column_prefers_the_volume_at_time(tmp_path):
    results = _results()
    results["rows"][0]["rel_volume_time"] = 1.15          # ENI : 2,4 sur 10 j, 1,15 à l'heure
    ws = load_workbook(str(xo.write_workbook(results, tmp_path)))["Candidats"]
    assert ws.cell(row=xo.DATA_START_ROW, column=8).value == 1.15
    assert ws.cell(row=xo.DATA_START_ROW, column=8).number_format == xo.FMT_RATIO
    # volume à l'heure inconnu (univers de repli, colonne morte) : le 10 j, comme avant
    results["rows"][0]["rel_volume_time"] = None
    ws = load_workbook(str(xo.write_workbook(results, tmp_path)))["Candidats"]
    assert ws.cell(row=xo.DATA_START_ROW, column=8).value == 2.4


def test_empty_results_still_write(tmp_path):
    path = xo.write_workbook({"generated_at": "2026-10-17T07:00:00+00:00", "rows": []}, tmp_path, "fr")
    wb = load_workbook(str(path))
    assert wb["Candidats"].cell(row=xo.DATA_START_ROW, column=1).value is None
