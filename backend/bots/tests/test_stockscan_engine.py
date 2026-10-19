"""Scanner d'actions — le scan de bout en bout, tout injecté (zéro réseau).

Univers : la fixture SYNTHÉTIQUE du screener (Italie). Signaux du coach :
chargeurs factices. News fraîches : fonction factice. Horloge épinglée au
2026-10-17 07:00 UTC (ENI publie ses résultats le 24/10, dans la fenêtre de 7 j ;
UniCredit le 29/10 et Ferrari le 04/11, hors fenêtre). En fin de fichier, une
régression sur la capture RÉELLE du 19/10 à 13:30 UTC (5 marchés).
"""
import json
import os
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from backend.bots.stockscan import __main__ as cli
from backend.bots.stockscan import coach as coachmod
from backend.bots.stockscan import config as cfgmod
from backend.bots.stockscan import engine
from backend.bots.stockscan import excel_out
from backend.bots.stockscan import markets as mk
from backend.bots.stockscan import screener as sc

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
NOW = datetime(2026, 10, 17, 7, 0, tzinfo=timezone.utc)


def _iso(hours_ago):
    return (NOW - timedelta(hours=hours_ago)).isoformat()


def _screener_payload():
    with open(os.path.join(FIXTURES, "stockscan", "screener_italy_synthetic.json"), encoding="utf-8") as f:
        return json.load(f)


class FakeScreener:
    def __init__(self, fail=()):
        self.fail = set(fail)
        self.calls = []

    def scan(self, market, cfg):
        self.calls.append(market.id)
        if market.id in self.fail or market.id != "italy":
            raise sc.ScreenerError("HTTP 403", 403)
        return _screener_payload(), {"status": 200, "attempts": 1, "ignored_unknown": False}


class FakeLoaders(coachmod.Loaders):
    def __init__(self, events=None, hypotheses=None, watchlist=None, tv_state=None):
        self._events = events or []
        self._hyps = hypotheses or []
        self._watch = watchlist or []
        self._tv = tv_state or {}

    def users(self):
        return ["massii"]

    def events(self, username):
        return list(self._events)

    def trends(self, now):
        return {}

    def tv_state(self):
        return self._tv

    def hypotheses(self):
        return list(self._hyps)

    def portfolio(self, username):
        return {"positions": []}

    def watchlist(self, username):
        return list(self._watch)


def _tv_item(title, minutes_ago=60, provider="reuters"):
    return {"id": title, "title": title, "provider": provider,
            "published": (NOW - timedelta(minutes=minutes_ago)).isoformat(),
            "url": "https://www.tradingview.com/news/" + title.replace(" ", "-")}


def _news_fetch(script, calls=None):
    def fetch(tv_symbol, lang):
        if calls is not None:
            calls.append((tv_symbol, lang))
        return script.get((tv_symbol, lang), ([], 200))
    return fetch


DEFAULT_LOADERS = dict(
    events=[{"ts": _iso(2), "symbol": "ENI.MI", "title": "Eni beats estimates on trading",
             "link": "https://n/eni", "sentiment": "pos", "src": "pressefi"}],
    hypotheses=[{"status": "open", "tickers": ["LDO.MI"], "direction": "up",
                 "thesis": "Commandes défense", "confidence": "moyenne"}],
    watchlist=[{"symbol": "UCG.MI", "name": "UniCredit"}],
    tv_state={"items_by_symbol": {"MIL:STLAM": {"fetched_at": _iso(0.1), "items": [
        {"id": "s1", "title": "Stellantis issues profit warning", "published": _iso(3),
         "provider": "reuters", "url": "https://www.tradingview.com/news/s1",
         "symbol": "STLAM.MI", "tv_symbol": "MIL:STLAM"}]}}},
)

NEWS_SCRIPT = {
    ("MIL:ENI", "en"): ([_tv_item("Eni signs gas supply deal with Algeria")], 200),
    ("MIL:UCG", "en"): ([_tv_item("UniCredit opens new branch in Milan", 120)], 200),
}


def _deps(screener=None, loaders=None, news=None, **kw):
    return engine.Deps(screener=screener or FakeScreener(),
                       coach_loaders=loaders or FakeLoaders(**DEFAULT_LOADERS),
                       news_fetch=news if news is not None else _news_fetch(NEWS_SCRIPT),
                       now=lambda: NOW, sleep=lambda s: None, monotonic=lambda: 0.0, **kw)


def _cfg(**over):
    base = {"markets": ["italy", "uk"], "min_value_meur": 5, "shortlist": 40, "top_n": 25}
    base.update(over)
    return cfgmod.validate(base)[0]


def _run(tmp_path, cfg=None, deps=None, **kw):
    logs = []
    results = engine.run_scan(cfg or _cfg(), "massii", tmp_path, deps or _deps(),
                              log=logs.append, **kw)
    return results, logs


def test_full_scan_ranks_catalysts_and_writes_files(tmp_path):
    results, logs = _run(tmp_path)
    symbols = [r["symbol"] for r in results["rows"]]
    # ENI : deux news positives/neutres + résultats proches + volume ; LDO : hypothèse
    # du coach + gros volume mais « déjà parti » ; UCG : une news neutre ; STLAM :
    # alerte rouge. RACE et ENEL n'ont aucun catalyseur ; TINY est trop peu liquide.
    assert symbols == ["ENI.MI", "LDO.MI", "UCG.MI", "STLAM.MI"]
    assert [r["rank"] for r in results["rows"]] == [1, 2, 3, 4]

    eni = results["rows"][0]
    assert "earnings_soon" in eni["flags"]
    # Samedi : Milan a clôturé vendredi, les news d'ENI sont postérieures ->
    # le volume et la variation de vendredi ne les confirment pas (sessions.py).
    assert {"positive", "earnings", "sources"} <= set(eni["components"])
    assert "volume" not in eni["components"] and "aligned" not in eni["components"]
    assert "awaiting_open" in eni["flags"]
    assert eni["session"]["state"] == "closed"
    assert eni["session"]["last_close"] == "2026-10-16T15:30:00+00:00"
    # l'état vient de TradingView (hors séance, barre du vendredi), pas de la table
    assert eni["session"]["source"] == "tradingview"
    assert eni["session"]["date"] == "2026-10-16" and eni["session"]["stale"] is True
    assert results["sessions"]["italy"] == {"state": "closed", "date": "2026-10-16",
                                            "stale": True, "elapsed": None,
                                            "source": "tradingview"}
    assert eni["links"]["chart"] == "https://www.tradingview.com/chart/?symbol=MIL%3AENI"
    assert len(eni["news"]) == 2

    ldo = results["rows"][1]
    assert "late" in ldo["flags"] and ldo["components"]["coach_up"] == 10
    assert ldo["coach"]["hypotheses"][0]["thesis"] == "Commandes défense"

    stl = results["rows"][3]
    assert "red" in stl["flags"] and stl["score"] == 0
    assert stl["news"][0]["sentiment"] == "neg"

    ucg = results["rows"][2]
    assert "followed" in ucg["flags"]

    c = results["counts"]
    assert c["universe"] == 6 and c["shortlist"] == 6 and c["kept"] == 4
    assert c["excluded_no_catalyst"] == 2 and c["excluded_red"] == 0
    # par marché : médiane du volume (6 titres < 8 : aucune) et présélectionnés
    status = {m["id"]: m for m in results["markets"]}
    assert status["italy"]["volume_median"] is None and status["italy"]["shortlisted"] == 6
    assert status["uk"]["volume_median"] is None and status["uk"]["shortlisted"] == 0
    assert "Présélection : 6 titres — italy 6" in logs
    assert c["news_from_cache"] == 1        # STLAM : cache du coach frais, pas de requête
    assert results["source"] == "screener"
    assert {"code": "market_failed", "market": "uk", "error": "HTTP 403"} in results["warnings"]

    on_disk = json.loads((tmp_path / "results.json").read_text(encoding="utf-8"))
    assert on_disk["rows"] == results["rows"]
    assert json.loads((tmp_path / "progress.json").read_text(encoding="utf-8"))["phase"] == "done"
    assert results["excel"] and (tmp_path / results["excel"]).is_file()
    wb = load_workbook(str(tmp_path / results["excel"]))
    # une feuille par marché, en italien par défaut (comme le Bond Scanner)
    assert results["excel"].startswith("Opportunita_Azioni_")
    assert wb.sheetnames == ["Italia", "Regno Unito", "Info"]
    # chaque compteur du moteur a un nom lisible dans la feuille « Info »
    assert set(results["counts"]) <= set(excel_out.LABELS["it"]["info"]["count_names"])
    assert any("Coach :" in line for line in logs)


def test_exclude_red_removes_and_counts(tmp_path):
    results, _ = _run(tmp_path, cfg=_cfg(exclude_red=True))
    assert "STLAM.MI" not in [r["symbol"] for r in results["rows"]]
    assert results["counts"]["excluded_red"] == 1


def test_top_n_cuts_the_list(tmp_path):
    results, _ = _run(tmp_path, cfg=_cfg(top_n=5, shortlist=10))
    assert len(results["rows"]) == 4             # moins de candidats que top_n
    results, _ = _run(tmp_path, cfg=_cfg(top_n=5, shortlist=10, exclude_red=True))
    assert results["counts"]["kept"] == 3


def test_shortlist_limits_news_requests(tmp_path):
    calls = []
    deps = _deps(news=_news_fetch(NEWS_SCRIPT, calls))
    cfg = _cfg()
    cfg["shortlist"] = 2                         # sous le plancher de validate : forcé ici
    results, _ = _run(tmp_path, cfg=cfg, deps=deps)
    assert results["counts"]["shortlist"] == 2
    assert len(calls) <= 2


def test_previous_run_marks_new_rows(tmp_path):
    results, _ = _run(tmp_path, previous_symbols={"ENI.MI"})
    flags = {r["symbol"]: r["flags"] for r in results["rows"]}
    assert "new" not in flags["ENI.MI"]
    assert "new" in flags["LDO.MI"]
    results, _ = _run(tmp_path, previous_symbols=None)
    assert all("new" not in r["flags"] for r in results["rows"])


def test_empty_coach_is_a_warning_not_a_silence(tmp_path):
    results, _ = _run(tmp_path, deps=_deps(loaders=FakeLoaders()))
    assert {"code": "coach_empty"} in results["warnings"]


def test_news_breaker_marks_unverified(tmp_path):
    script = {("MIL:%s" % t, "en"): (None, 429) for t in ("ENI", "UCG", "LDO", "ENEL", "RACE")}
    results, _ = _run(tmp_path, deps=_deps(news=_news_fetch(script)))
    assert any(w["code"] == "news_breaker" for w in results["warnings"])
    eni = [r for r in results["rows"] if r["symbol"] == "ENI.MI"][0]
    assert "news_unverified" in eni["flags"]


def test_stop_flag_stops_before_writing_results(tmp_path):
    (tmp_path / "stop.flag").write_text("1", encoding="utf-8")
    with pytest.raises(engine.Stopped):
        _run(tmp_path)
    assert not (tmp_path / "results.json").exists()


def test_fallback_universe_when_screener_fails_everywhere(tmp_path):
    def chart(symbol):
        closes = [100.0, 101.0, 102.0, 101.5, 103.0, 104.0, 106.0]
        vols = [1e6, 1.1e6, 0.9e6, 1e6, 1e6, 1e6, 3e6]
        return {"chart": {"result": [{
            "meta": {"regularMarketPrice": 106.0, "currency": "EUR", "longName": "Eni S.p.A.",
                     "exchangeName": "MIL"},
            "timestamp": list(range(len(closes))),
            "indicators": {"quote": [{"open": closes, "high": closes, "low": closes,
                                      "close": closes, "volume": vols}]}}]}}

    deps = _deps(screener=FakeScreener(fail={"italy"}), chart=chart,
                 fallback_symbols=lambda u: ["ENI.MI", "AAPL", "NESN.SW"])
    results, _ = _run(tmp_path, cfg=_cfg(markets=["italy"]), deps=deps)
    assert results["source"] == "fallback"
    assert {"code": "fallback_universe"} in results["warnings"]
    assert results["counts"]["universe"] == 1          # seul ENI.MI est sur le marché choisi
    eni = results["rows"][0]
    assert eni["tv_symbol"] == "MIL:ENI"
    assert eni["rel_volume"] == 3.0
    # Yahoo ne dit rien de la séance : la table prend le relais (samedi : fermée)
    assert (eni["session"]["source"], eni["session"]["state"], eni["session"]["date"],
            eni["session"]["stale"]) == ("table", "closed", None, False)
    assert eni["session"]["last_close"] == "2026-10-16T15:30:00+00:00"
    assert eni["change_pct"] == round((106 - 104) / 104 * 100, 2)


def test_fallback_breaker_stops_after_three_failures(tmp_path):
    calls = []

    def chart(symbol):
        calls.append(symbol)
        raise OSError("network down")

    deps = _deps(screener=FakeScreener(fail={"america"}), chart=chart,
                 fallback_symbols=lambda u: ["AAPL", "MSFT", "NVDA", "TSLA", "AMZN", "META"])
    results, logs = _run(tmp_path, cfg=_cfg(markets=["america"]), deps=deps)
    assert calls == ["AAPL", "MSFT", "NVDA"]          # pas 6 × 13 s de vide
    assert {"code": "fallback_failed"} in results["warnings"]
    assert any("repli abandonné" in line for line in logs)


def test_universe_empty_everywhere_still_writes_results(tmp_path):
    deps = _deps(screener=FakeScreener(fail={"italy"}))
    results, _ = _run(tmp_path, cfg=_cfg(markets=["italy"]), deps=deps)
    assert results["rows"] == [] and results["counts"]["universe"] == 0
    assert (tmp_path / "results.json").is_file()


def test_row_from_chart_handles_pence_and_garbage():
    raw = {"chart": {"result": [{
        "meta": {"regularMarketPrice": 2450.0, "currency": "GBp", "shortName": "Shell"},
        "timestamp": [1, 2, 3],
        "indicators": {"quote": [{"open": [2400, 2420, 2450], "high": [2400, 2420, 2450],
                                  "low": [2400, 2420, 2450], "close": [2400, 2420, 2450],
                                  "volume": [1e7, 1e7, 1e7]}]}}]}}
    row = engine.row_from_chart("SHEL.L", raw)
    assert row["currency"] == "GBX"
    assert row["avg_value_meur"] == round(1e7 * 2450 * 0.0117 / 1e6, 2)
    assert row["tv_symbol"] == "LSE:SHEL"
    assert engine.row_from_chart("SHEL.L", {"chart": {"result": []}}) is None
    assert engine.row_from_chart("XX.ZZ", raw) is None          # place inconnue
    assert engine.row_from_chart("SHEL.L", None) is None


def test_yahoo_market_helpers():
    assert mk.market_of_yahoo("ENI.MI") == "italy"
    assert mk.market_of_yahoo("BRK-B") == "america"
    assert mk.market_of_yahoo("EURUSD=X") is None
    assert mk.yahoo_to_tv("RO.SW") == "SIX:ROG"
    assert mk.yahoo_to_tv("BRK-B") == "NASDAQ:BRK.B"
    assert mk.yahoo_to_tv("ASML.AS") == "EURONEXT:ASML"


# --------------------------------------------------------------------------- #
# __main__
# --------------------------------------------------------------------------- #

def test_main_reads_config_and_returns_codes(tmp_path):
    run_dir = tmp_path / "runs" / ("a" * 32)
    run_dir.mkdir(parents=True)
    (run_dir / "config.json").write_text(json.dumps({"markets": ["italy"], "lang": "it"}),
                                         encoding="utf-8")
    code = cli.main(["--out", str(run_dir), "--config", str(run_dir / "config.json"),
                     "--user", "massii"], deps=_deps())
    assert code == 0
    results = json.loads((run_dir / "results.json").read_text(encoding="utf-8"))
    assert results["config"]["lang"] == "it"
    assert results["excel"].startswith("Opportunita_Azioni_")


def test_main_previous_run_feeds_new_flag(tmp_path):
    runs = tmp_path / "runs"
    old = runs / ("b" * 32)
    old.mkdir(parents=True)
    (old / "results.json").write_text(json.dumps({"rows": [{"symbol": "ENI.MI"}]}), encoding="utf-8")
    assert cli.previous_symbols(runs / ("c" * 32)) == {"ENI.MI"}


def test_main_stop_and_crash_codes(tmp_path):
    run_dir = tmp_path / "r"
    run_dir.mkdir()
    (run_dir / "stop.flag").write_text("1", encoding="utf-8")
    assert cli.main(["--out", str(run_dir)], deps=_deps()) == 3

    class Boom(FakeScreener):
        def scan(self, market, cfg):
            raise RuntimeError("bug")

    run2 = tmp_path / "r2"
    assert cli.main(["--out", str(run2), "--config", str(tmp_path / "absent.json")],
                    deps=_deps(screener=Boom())) == 1


def test_main_empty_universe_returns_2(tmp_path):
    run_dir = tmp_path / "r3"
    code = cli.main(["--out", str(run_dir)], deps=_deps(screener=FakeScreener(fail={"italy"})))
    assert code == 2
    assert (run_dir / "results.json").is_file()


def test_probe_writes_raw_responses(tmp_path):
    class Resp:
        status_code = 200

        def json(self):
            with open(os.path.join(FIXTURES, "tvcoach", "tv_news_aapl.json"), encoding="utf-8") as f:
                return json.load(f)

    class News:
        def get(self, url, headers=None):
            return Resp()

    code = cli.probe(str(tmp_path), "italy", None, screener=FakeScreener(), news_get=News())
    assert code == 0
    assert (tmp_path / "screener_italy.json").is_file()
    assert (tmp_path / "tvnews_MIL_UCG.json").is_file()


def test_probe_reports_screener_failure(tmp_path):
    assert cli.probe(str(tmp_path), "italy", None, screener=FakeScreener(fail={"italy"})) == 1
    assert cli.probe(str(tmp_path), "japan", None, screener=FakeScreener()) == 1


class SessionScreener(FakeScreener):
    """La fixture italienne avec l'état de séance que TradingView rendrait à un
    autre instant : ``current_session`` et ouverture de la barre du jour."""

    def __init__(self, current_session, bar_open, **kw):
        super().__init__(**kw)
        self.current_session, self.bar_open = current_session, bar_open

    def scan(self, market, cfg):
        payload, meta = super().scan(market, cfg)
        i_session, i_time = sc.COLUMNS.index("current_session"), sc.COLUMNS.index("time")
        for raw in payload["data"]:
            if raw["s"]:
                raw["d"][i_session] = self.current_session
                raw["d"][i_time] = int(self.bar_open.timestamp())
        return payload, meta


def test_open_session_volume_is_no_longer_scaled_by_elapsed_time(tmp_path):
    # Mercredi 21/10 13:15 UTC = 15:15 à Milan, séance ouverte. Le volume n'est
    # plus divisé par la part de séance écoulée (courbe en U : 74 % du temps
    # n'est pas 74 % du volume) : c'est le volume « à l'heure », rapporté à la
    # médiane de son marché — ici 6 titres, trop peu (8) : diviseur 1.
    wednesday = datetime(2026, 10, 21, 13, 15, tzinfo=timezone.utc)
    deps = _deps(screener=SessionScreener(
        "market", datetime(2026, 10, 21, 7, 0, tzinfo=timezone.utc)))
    deps.now = lambda: wednesday
    results = engine.run_scan(_cfg(), "massii", tmp_path, deps, log=lambda m: None)
    assert not {w["code"] for w in results["warnings"]} & {"volume_adjusted", "session_uncertain"}
    assert results["markets"][0]["volume_median"] is None
    assert results["rows"]
    session = results["rows"][0]["session"]
    assert (session["source"], session["state"], session["date"], session["stale"]) == \
        ("tradingview", "open", "2026-10-21", False)
    assert session["elapsed"] == round(6.25 / 8.5, 3) and session["last_close"] is None
    for row in results["rows"]:
        assert row["rel_volume_raw"] == row["rel_volume_time"]
        assert row["rel_volume"] == round(row["rel_volume_raw"], 2)


# --------------------------------------------------------------------------- #
# Régression sur la capture RÉELLE du screener (lundi 19/10/2026 13:30 UTC)
# --------------------------------------------------------------------------- #

SNAPSHOT_NOW = datetime(2026, 10, 19, 13, 30, 37, tzinfo=timezone.utc)
SNAPSHOT_MARKETS = ["switzerland", "italy", "germany", "france", "america"]


class SnapshotScreener:
    """Sert la capture réelle, marché par marché (Europe en séance, US en
    pré-ouverture avec les chiffres du vendredi 16/10). ``width`` tronque
    chaque ligne : 25 = les colonnes d'avant le 19/10, comme si TradingView
    avait renommé les quatre nouvelles."""

    def __init__(self, width=None):
        with open(os.path.join(FIXTURES, "stockscan", "screener_snapshot_2026-10-19T1330Z.json"),
                  encoding="utf-8") as f:
            self.snapshot = json.load(f)
        self.width = width

    def scan(self, market, cfg):
        block = self.snapshot["markets"].get(market.id)
        if block is None:
            raise sc.ScreenerError("HTTP 404", 404)
        data = [{"s": r["s"], "d": r["d"][:self.width] if self.width else r["d"]}
                for r in block["data"]]
        return ({"totalCount": block["totalCount"], "data": data},
                {"status": 200, "attempts": 1, "ignored_unknown": bool(self.width)})


def _snapshot_deps(screener):
    return engine.Deps(screener=screener, coach_loaders=FakeLoaders(), news_fetch=None,
                       now=lambda: SNAPSHOT_NOW, sleep=lambda s: None, monotonic=lambda: 0.0)


def test_real_snapshot_every_market_reaches_the_shortlist(tmp_path):
    # Sans coach ni news, seuls volume, variation et résultats proches classent.
    # Volume « à l'heure » ÷ médiane de son marché + quota de 50 % : chaque
    # place est présélectionnée, pas seulement les États-Unis. Chaque ligne
    # garde ses deux volumes — un champ oublié dans une couche serait une
    # fonctionnalité morte sans aucune erreur (piège #61 du dépôt).
    cfg = _cfg(markets=SNAPSHOT_MARKETS, shortlist=20, top_n=10)
    results, logs = _run(tmp_path, cfg=cfg, deps=_snapshot_deps(SnapshotScreener()))

    status = {m["id"]: m for m in results["markets"]}
    assert [m["id"] for m in results["markets"]] == SNAPSHOT_MARKETS
    assert {m: status[m]["volume_median"] for m in SNAPSHOT_MARKETS} == {
        "switzerland": 1.064, "italy": 0.912, "germany": 0.856, "france": 0.987,
        "america": 1.187}
    shortlisted = {m: status[m]["shortlisted"] for m in SNAPSHOT_MARKETS}
    assert all(n >= 2 for n in shortlisted.values())
    assert sum(shortlisted.values()) == results["counts"]["shortlist"] == 20
    assert shortlisted["america"] < 20
    assert not {w["code"] for w in results["warnings"]} & {"volume_adjusted", "session_uncertain",
                                                           "degraded_columns"}
    # jamais de dict dans les compteurs : l'onglet Infos de l'Excel écrit
    # chaque compteur dans UNE cellule
    assert all(isinstance(v, int) for v in results["counts"].values())

    assert results["rows"]                     # une liste vide prouverait tout
    for row in results["rows"]:
        assert isinstance(row["rel_volume_raw"], float), row["symbol"]
        assert row["rel_volume_time"] == row["rel_volume_raw"], row["symbol"]
    on_disk = json.loads((tmp_path / "results.json").read_text(encoding="utf-8"))
    assert on_disk["rows"] == results["rows"]

    assert ("Volume relatif médian à l'heure : switzerland 1,06 · italy 0,91 · "
            "germany 0,86 · france 0,99 · america 1,19") in logs
    expected = ", ".join("%s %d" % (m, n) for m, n in
                         sorted(shortlisted.items(), key=lambda kv: (-kv[1], kv[0])))
    assert "Présélection : 20 titres — " + expected in logs


def test_dead_volume_at_time_column_falls_back_to_the_ten_day_volume(tmp_path):
    # TradingView renomme les quatre nouvelles colonnes : le relevé rejoué les
    # rend vides. Le scan va au bout, le dit (degraded_columns), et le volume
    # retombe sur le 10 j — rapporté à sa médiane, bornée : en pleine séance
    # européenne (0,27 à 0,40) le diviseur plafonne à 0,5. Le journal ne doit
    # pas appeler « à l'heure » un volume qui ne l'est pas.
    cfg = _cfg(markets=SNAPSHOT_MARKETS, shortlist=20, top_n=10)
    results, logs = _run(tmp_path, cfg=cfg, deps=_snapshot_deps(SnapshotScreener(width=25)))

    assert {"code": "degraded_columns",
            "columns": ["current_session", "relative_volume_intraday|5", "time", "timezone"]} \
        in results["warnings"]
    assert {m["id"]: m["volume_median"] for m in results["markets"]} == {
        "switzerland": 0.37, "italy": 0.403, "germany": 0.305, "france": 0.271,
        "america": 1.223}
    assert results["rows"]
    for row in results["rows"]:
        assert row["rel_volume_time"] is None
        divisor = 1.223 if row["market"] == "america" else 0.5
        assert abs(row["rel_volume"] - row["rel_volume_raw"] / divisor) <= 0.006, row["symbol"]
    assert ("Volume relatif médian à l'heure : switzerland 0,37 (10 j) · italy 0,40 (10 j) · "
            "germany 0,30 (10 j) · france 0,27 (10 j) · america 1,22 (10 j)") in logs


def test_real_snapshot_sessions_come_from_tradingview(tmp_path):
    # 13:30:37 UTC = 09:30 à New York : la table dirait la bourse ouverte, mais
    # les données (différées) sont en PRÉ-OUVERTURE, barre du vendredi 16/10 —
    # des chiffres d'avant les news du week-end, qui ne peuvent ni les
    # confirmer ni les dire « déjà dans le cours » : « attend l'ouverture ».
    # Milan est en séance avec la barre du jour : cours vivants, rien n'attend.
    # Des news du coach font passer un titre US et un titre italien SANS
    # résultats proches le filtre « catalyseur ».
    news = {"NVDA": ["Nvidia signs a data-center supply deal",
                     "Nvidia opens a research lab in Tel Aviv"],
            "UCG.MI": ["UniCredit completes the sale of its Russian unit",
                       "UniCredit names a new head of retail banking"]}
    events = [{"ts": (SNAPSHOT_NOW - timedelta(hours=2 + i)).isoformat(), "symbol": symbol,
               "title": title, "link": "https://n/%s/%d" % (symbol, i),
               "sentiment": "pos", "src": "pressefi"}
              for symbol, titles in news.items() for i, title in enumerate(titles)]
    deps = _snapshot_deps(SnapshotScreener())
    deps.coach_loaders = FakeLoaders(events=events)
    cfg = _cfg(markets=SNAPSHOT_MARKETS, shortlist=40, top_n=25)
    results, _ = _run(tmp_path, cfg=cfg, deps=deps)
    rows = {r["symbol"]: r for r in results["rows"]}

    nvda = rows["NVDA"]
    assert nvda["session"] == {"source": "tradingview", "state": "pre_open",
                               "date": "2026-10-16", "stale": True, "elapsed": None,
                               "local_time": "09:30",
                               "last_close": "2026-10-16T20:00:00+00:00"}
    assert "awaiting_open" in nvda["flags"]
    assert not {"volume", "aligned"} & set(nvda["components"])

    ucg = rows["UCG.MI"]
    assert (ucg["session"]["source"], ucg["session"]["state"], ucg["session"]["date"],
            ucg["session"]["stale"], ucg["session"]["last_close"]) == \
        ("tradingview", "open", "2026-10-19", False, None)
    assert ucg["session"]["elapsed"] is not None
    assert "awaiting_open" not in ucg["flags"]           # cours vivants : rien n'attend

    assert results["sessions"]["america"] == {"state": "pre_open", "date": "2026-10-16",
                                              "stale": True, "elapsed": None,
                                              "source": "tradingview"}
    assert {m: (s["state"], s["date"], s["stale"]) for m, s in results["sessions"].items()
            if m != "america"} == {m: ("open", "2026-10-19", False)
                                   for m in SNAPSHOT_MARKETS if m != "america"}
